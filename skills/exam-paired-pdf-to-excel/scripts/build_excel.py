import pickle, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

q,a,chap_titles=pickle.load(open('parsed.pkl','rb'))

def clean(s):
    if s is None: return ''
    s=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]','',s)  # drop control chars
    return s.strip()

def chap_cn(n):
    cn=['零','一','二','三','四','五','六','七','八','九','十']
    return cn[n] if n<=10 else str(n)

# validate answers
bad=[]
for k in a:
    ans=a[k]['ans']
    if not re.fullmatch(r'[A-E]+', ans):
        bad.append((k,ans,'invalid letters'))
    elif q[k]['sec']=='单选' and len(ans)!=1:
        bad.append((k,ans,'单选 but multi-letter'))
    elif q[k]['sec']=='多选' and len(ans)<2:
        bad.append((k,ans,'多选 but single-letter'))
print('answer anomalies:',bad)

# build rows sorted by chapter then qnum
keys=sorted(q.keys(), key=lambda k:(k[0],k[1]))
wb=Workbook()
ws=wb.active
ws.title='史纲720题'
headers=['ID','题目','题型','分数','难度','选项A','选项B','选项C','选项D','选项E','答案','解析','一级目录','二级目录','章节']
ws.append(headers)

first_dir='27考研政治·米鹏720题（中国近现代史纲要）'
row=0
for k in keys:
    row+=1
    ch,qn=k
    sec=q[k]['sec']         # 单选/多选
    题型='单选题' if sec=='单选' else '多选题'
    二级='单项选择题' if sec=='单选' else '多项选择题'
    title=clean(chap_titles.get(ch,''))
    章节=f'第{chap_cn(ch)}章 {title}'
    opts={L:T for L,T in q[k]['options']}
    rowdata=[
        row,
        clean(q[k]['text']),
        题型,
        '',  # 分数
        '',  # 难度
        clean(opts.get('A','')),
        clean(opts.get('B','')),
        clean(opts.get('C','')),
        clean(opts.get('D','')),
        clean(opts.get('E','')),
        a[k]['ans'],
        clean(a[k]['content']),
        first_dir,
        二级,
        章节,
    ]
    ws.append(rowdata)

# styling
hdr_fill=PatternFill('solid',fgColor='FF07C160')
hdr_font=Font(bold=True,color='FFFFFFFF')
thin=Side(style='thin',color='FFD0D0D0')
border=Border(left=thin,right=thin,top=thin,bottom=thin)
for c in range(1,len(headers)+1):
    cell=ws.cell(row=1,column=c)
    cell.fill=hdr_fill; cell.font=hdr_font
    cell.alignment=Alignment(horizontal='center',vertical='center')
    cell.border=border
widths=[5,60,8,6,6,28,28,28,28,28,8,80,26,14,30]
for i,w in enumerate(widths,1):
    ws.column_dimensions[chr(64+i) if i<=26 else 'A'].width=w
# wrap text for 题目/解析/选项
for r in range(2,ws.max_row+1):
    for c in [2,6,7,8,9,10,12]:
        ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical='top')
    ws.cell(row=r,column=11).alignment=Alignment(horizontal='center',vertical='center')
    for c in range(1,len(headers)+1):
        ws.cell(row=r,column=c).border=border
ws.freeze_panes='A2'
wb.save('720-史纲题库.xlsx')
print('SAVED 720-史纲题库.xlsx  rows=',ws.max_row-1)
